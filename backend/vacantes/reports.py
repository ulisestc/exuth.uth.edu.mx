import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from django.utils import timezone
from .models import Colocacion, Postulacion, Vacante
from profiles.models import PadronEgresado, Empresa
from core.models import Carrera

def _apply_uth_header(sheet, title, subtitle, max_col):
    """Aplica la cabecera institucional oficial UTH al inicio de una hoja Excel."""
    fill_dark = PatternFill(start_color="2D2926", end_color="2D2926", fill_type="solid")
    fill_green = PatternFill(start_color="00A887", end_color="00A887", fill_type="solid")
    
    font_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=9, bold=True, color="C2BA98")
    font_banner = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_meta = Font(name="Calibri", size=9, italic=True, color="636569")
    
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Fila 1: Título Institucional
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell1 = sheet.cell(row=1, column=1, value="UNIVERSIDAD TECNOLÓGICA DE HUEJOTZINGO")
    cell1.fill = fill_dark
    cell1.font = font_title
    cell1.alignment = align_center
    sheet.row_dimensions[1].height = 26
    
    # Fila 2: Dirección y Sistema
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    cell2 = sheet.cell(row=2, column=1, value="DIRECCIÓN DE VINCULACIÓN Y EXTENSIÓN UNIVERSITARIA · BOLSA DE TRABAJO EXUTH")
    cell2.fill = fill_dark
    cell2.font = font_subtitle
    cell2.alignment = align_center
    sheet.row_dimensions[2].height = 18
    
    # Fila 3: Banner del Reporte
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
    cell3 = sheet.cell(row=3, column=1, value=title.upper())
    cell3.fill = fill_green
    cell3.font = font_banner
    cell3.alignment = align_center
    sheet.row_dimensions[3].height = 24
    
    # Fila 4: Metadatos y Fecha de Corte
    fecha_emision = timezone.now().strftime("%d/%m/%Y %H:%M hrs")
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max_col)
    cell4 = sheet.cell(row=4, column=1, value=f"{subtitle} | Fecha y hora de emisión: {fecha_emision}")
    cell4.font = font_meta
    cell4.alignment = align_center
    sheet.row_dimensions[4].height = 18
    
    # Fila 5: Espaciador
    sheet.row_dimensions[5].height = 10

def _format_table_headers(sheet, headers, start_row=6):
    """Estiliza los encabezados de tabla con color Verde UTH y texto blanco en negrita."""
    fill_header = PatternFill(start_color="00A887", end_color="00A887", fill_type="solid")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style="thin", color="008F73"),
        right=Side(style="thin", color="008F73"),
        top=Side(style="thin", color="008F73"),
        bottom=Side(style="medium", color="005F4D")
    )
    
    sheet.row_dimensions[start_row].height = 28
    for col_idx, header_text in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col_idx, value=header_text)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

def _auto_adjust_columns(sheet, min_row=6, max_col=22):
    """Ajusta automáticamente el ancho de las columnas basándose en el contenido."""
    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(min_row, sheet.max_row + 1):
            val = sheet.cell(row=row, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        width = max(max_len + 4, 12)
        sheet.column_dimensions[col_letter].width = min(width, 45)

def generate_colocacion_excel_report():
    """
    Genera el Reporte Consolidado de Colocación Laboral en formato Excel (.xlsx).
    Cumple con los requerimientos de auditoría y acreditación para CACEI, CONAIC e ISO.
    """
    wb = openpyxl.Workbook()
    
    border_data = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )
    fill_zebra_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    font_data = Font(name="Calibri", size=9.5, color="1F2937")
    font_bold = Font(name="Calibri", size=9.5, bold=True, color="111827")
    
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    # =========================================================================
    # HOJA 1: COLOCACIONES LABORALES (CACEI / CONAIC / ISO)
    # =========================================================================
    ws_coloc = wb.active
    ws_coloc.title = "Colocaciones_CACEI_ISO"
    ws_coloc.views.sheetView[0].showGridLines = True
    
    headers_coloc = [
        "Folio",
        "Fecha de Colocación",
        "Matrícula UTH",
        "Nombre del Egresado",
        "CURP",
        "Género",
        "Teléfono Móvil",
        "Correo Electrónico",
        "Carrera / Programa Educativo",
        "Nivel Académico",
        "Periodo de Egreso (Padrón)",
        "Año Egreso",
        "Estatus Titulación",
        "Empresa Contratante",
        "Giro Empresarial",
        "Sector Industrial",
        "Clave Vacante",
        "Puesto Laboral",
        "Modalidad",
        "Rango Salarial",
        "Registrado Por",
        "Observaciones Institucionales"
    ]
    
    colocaciones = Colocacion.objects.select_related(
        'egresado__user',
        'egresado__carrera',
        'vacante__empresa__giro',
        'vacante__empresa__sector',
        'registrado_por'
    ).order_by('-fecha_colocacion', '-id')
    
    matriculas_coloc = [c.egresado.matricula for c in colocaciones if c.egresado]
    padron_map = {}
    if matriculas_coloc:
        for p in PadronEgresado.objects.filter(matricula__in=matriculas_coloc):
            padron_map[p.matricula] = p
            
    _apply_uth_header(
        ws_coloc,
        "Reporte Oficial de Colocación Laboral y Seguimiento de Egresados",
        f"Acreditación CACEI / CONAIC / ISO · Total de Contrataciones Concretadas: {colocaciones.count()}",
        len(headers_coloc)
    )
    _format_table_headers(ws_coloc, headers_coloc, start_row=6)
    
    current_row = 7
    if colocaciones.exists():
        for idx, col in enumerate(colocaciones, start=1):
            egr = col.egresado
            usr = egr.user if egr else None
            vac = col.vacante
            emp = vac.empresa if vac else None
            pad = padron_map.get(egr.matricula) if egr else None
            
            sueldo_str = "No especificado"
            if vac:
                if vac.salario_a_tratar:
                    sueldo_str = "A convenir en entrevista"
                elif vac.sueldo_minimo and vac.sueldo_maximo:
                    sueldo_str = f"${vac.sueldo_minimo:,.2f} - ${vac.sueldo_maximo:,.2f} MXN"
                elif vac.sueldo_minimo:
                    sueldo_str = f"Desde ${vac.sueldo_minimo:,.2f} MXN"
                    
            row_fill = fill_zebra_even if idx % 2 == 0 else fill_zebra_odd
            ws_coloc.row_dimensions[current_row].height = 20
            
            row_values = [
                f"COL-{col.id:04d}",
                col.fecha_colocacion.strftime("%d/%m/%Y") if col.fecha_colocacion else "N/A",
                egr.matricula if egr else "N/A",
                f"{usr.nombres} {usr.apellido_paterno} {usr.apellido_materno}".strip() if usr else "N/A",
                egr.curp if egr else "N/A",
                egr.get_genero_display() if egr else "N/A",
                egr.telefono_celular or (pad.telefono_movil if pad else "N/A"),
                usr.email if usr else "N/A",
                egr.carrera.nombre if (egr and egr.carrera) else "N/A",
                egr.get_nivel_estudios_display() if egr else "N/A",
                pad.periodo if pad and pad.periodo else "N/D",
                pad.anio_egreso if pad and pad.anio_egreso else "N/D",
                pad.estatus_titulacion if pad and pad.estatus_titulacion else "En Proceso",
                emp.nombre if emp else "Empresa Externa / No vinculada",
                emp.giro.nombre if (emp and emp.giro) else "N/A",
                emp.sector.nombre if (emp and emp.sector) else "N/A",
                vac.clave_vacante if vac else "N/A",
                vac.titulo if vac else "Puesto Externo",
                vac.get_modalidad_display() if vac else "N/A",
                sueldo_str,
                f"{col.registrado_por.nombres} ({col.registrado_por.rol})" if col.registrado_por else "Sistema UTH",
                col.observaciones or "Colocación validada satisfactoriamente."
            ]
            
            for c_idx, val in enumerate(row_values, start=1):
                c = ws_coloc.cell(row=current_row, column=c_idx, value=val)
                c.fill = row_fill
                c.border = border_data
                c.font = font_bold if c_idx in [1, 3, 14, 18] else font_data
                c.alignment = align_center if c_idx in [1, 2, 3, 5, 6, 7, 10, 11, 12, 17, 19] else align_left
                
            current_row += 1
    else:
        ws_coloc.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers_coloc))
        empty_cell = ws_coloc.cell(
            row=current_row, 
            column=1, 
            value="Actualmente no se han reportado contrataciones formalizadas en el sistema. Las colocaciones se generan de forma automática cuando una empresa empleadora acepta formalmente a un candidato turnado."
        )
        empty_cell.alignment = align_center
        empty_cell.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        ws_coloc.row_dimensions[current_row].height = 30
        current_row += 1

    ws_coloc.freeze_panes = "A7"
    _auto_adjust_columns(ws_coloc, min_row=6, max_col=len(headers_coloc))

    # =========================================================================
    # HOJA 2: AUDITORÍA DE POSTULACIONES Y FILTRO UTH
    # =========================================================================
    ws_posts = wb.create_sheet(title="Postulaciones_y_Filtro")
    ws_posts.views.sheetView[0].showGridLines = True
    
    headers_posts = [
        "ID Postulación",
        "Fecha de Registro",
        "Matrícula UTH",
        "Nombre del Egresado",
        "Carrera",
        "Clave Vacante",
        "Puesto Postulado",
        "Empresa",
        "Estatus de Vinculación",
        "CV Registrado",
        "Observaciones de Revisión UTH"
    ]
    
    postulaciones = Postulacion.objects.select_related(
        'egresado__user',
        'egresado__carrera',
        'vacante__empresa'
    ).order_by('-fecha_postulacion', '-id')
    
    _apply_uth_header(
        ws_posts,
        "Auditoría y Bitácora de Postulaciones Laborales",
        f"Seguimiento de Filtro Institucional · Total Postulaciones: {postulaciones.count()}",
        len(headers_posts)
    )
    _format_table_headers(ws_posts, headers_posts, start_row=6)
    
    row_posts = 7
    if postulaciones.exists():
        for idx, post in enumerate(postulaciones, start=1):
            egr = post.egresado
            usr = egr.user if egr else None
            vac = post.vacante
            emp = vac.empresa if vac else None
            
            row_fill = fill_zebra_even if idx % 2 == 0 else fill_zebra_odd
            ws_posts.row_dimensions[row_posts].height = 20
            
            post_values = [
                f"POST-{post.id:04d}",
                post.fecha_postulacion.strftime("%d/%m/%Y %H:%M") if post.fecha_postulacion else "N/A",
                egr.matricula if egr else "N/A",
                f"{usr.nombres} {usr.apellido_paterno} {usr.apellido_materno}".strip() if usr else "N/A",
                egr.carrera.nombre if (egr and egr.carrera) else "N/A",
                vac.clave_vacante if vac else "N/A",
                vac.titulo if vac else "N/A",
                emp.nombre if emp else "N/A",
                post.get_estado_display(),
                "Sí (PDF adjunto)" if egr and egr.cv else "Sin CV cargado",
                post.notas_uth or "En espera de validación de pertinencia."
            ]
            
            for c_idx, val in enumerate(post_values, start=1):
                c = ws_posts.cell(row=row_posts, column=c_idx, value=val)
                c.fill = row_fill
                c.border = border_data
                c.font = font_bold if c_idx in [1, 4, 7, 9] else font_data
                c.alignment = align_center if c_idx in [1, 2, 3, 6, 9, 10] else align_left
                
            row_posts += 1
    else:
        ws_posts.merge_cells(start_row=row_posts, start_column=1, end_row=row_posts, end_column=len(headers_posts))
        empty_cell = ws_posts.cell(
            row=row_posts, 
            column=1, 
            value="No existen postulaciones registradas en el sistema a la fecha de corte."
        )
        empty_cell.alignment = align_center
        empty_cell.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        ws_posts.row_dimensions[row_posts].height = 30

    ws_posts.freeze_panes = "A7"
    _auto_adjust_columns(ws_posts, min_row=6, max_col=len(headers_posts))

    # =========================================================================
    # HOJA 3: INDICADORES ESTADÍSTICOS Y KPIS (CACEI / CONAIC)
    # =========================================================================
    ws_stats = wb.create_sheet(title="Estadisticas_y_KPIs")
    ws_stats.views.sheetView[0].showGridLines = True
    
    _apply_uth_header(
        ws_stats,
        "Indicadores Institucionales de Efectividad y Empleabilidad",
        "Métricas cuantitativas para acreditación de programas educativos",
        6
    )
    
    fill_subhead = PatternFill(start_color="2D2926", end_color="2D2926", fill_type="solid")
    font_subhead = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    
    cell_kpi = ws_stats.cell(row=6, column=1, value="INDICADOR INSTITUCIONAL")
    cell_kpi.fill = fill_subhead
    cell_kpi.font = font_subhead
    cell_kpi.alignment = align_left
    
    cell_kpi_val = ws_stats.cell(row=6, column=2, value="VALOR")
    cell_kpi_val.fill = fill_subhead
    cell_kpi_val.font = font_subhead
    cell_kpi_val.alignment = align_center
    
    total_posts = postulaciones.count()
    total_coloc = colocaciones.count()
    turnadas_count = postulaciones.filter(estado__in=['enviada_empresa', 'Aceptada']).count()
    tasa_efectividad = f"{(total_coloc / turnadas_count * 100):.1f}%" if turnadas_count > 0 else "0.0%"
    empresas_activas = Empresa.objects.filter(status='aprobada').count()
    vacantes_activas = Vacante.objects.filter(status='aprobada').count()
    
    kpis = [
        ("Total de Egresados Colocados Formalmente", total_coloc),
        ("Total de Postulaciones Registradas", total_posts),
        ("Postulaciones Turnadas a Empresas", turnadas_count),
        ("Tasa de Efectividad de Colocación (Colocados / Turnados)", tasa_efectividad),
        ("Empresas Vinculadas Validadas", empresas_activas),
        ("Ofertas de Empleo Aprobadas Vigentes", vacantes_activas),
    ]
    
    for r_idx, (nombre_kpi, val_kpi) in enumerate(kpis, start=7):
        ws_stats.row_dimensions[r_idx].height = 20
        c1 = ws_stats.cell(row=r_idx, column=1, value=nombre_kpi)
        c2 = ws_stats.cell(row=r_idx, column=2, value=str(val_kpi))
        c1.fill = fill_zebra_odd
        c2.fill = fill_zebra_odd
        c1.border = border_data
        c2.border = border_data
        c1.font = font_data
        c2.font = font_bold
        c2.alignment = align_center

    ws_stats.merge_cells("D6:F6")
    cell_carr = ws_stats.cell(row=6, column=4, value="DISTRIBUCIÓN DE COLOCACIÓN POR PROGRAMA EDUCATIVO")
    cell_carr.fill = PatternFill(start_color="00A887", end_color="00A887", fill_type="solid")
    cell_carr.font = font_subhead
    cell_carr.alignment = align_left
    
    headers_carr = ["Carrera / Programa Educativo", "Colocados", "% del Total"]
    for c_i, h_txt in enumerate(headers_carr, start=4):
        c = ws_stats.cell(row=7, column=c_i, value=h_txt)
        c.fill = fill_subhead
        c.font = font_subhead
        c.border = border_data
        c.alignment = align_center if c_i > 4 else align_left
        
    carreras = Carrera.objects.all().order_by('nombre')
    r_carr = 8
    for car in carreras:
        coloc_en_car = colocaciones.filter(egresado__carrera=car).count()
        porcentaje = f"{(coloc_en_car / total_coloc * 100):.1f}%" if total_coloc > 0 else "0.0%"
        ws_stats.row_dimensions[r_carr].height = 19
        
        c_nom = ws_stats.cell(row=r_carr, column=4, value=car.nombre)
        c_cnt = ws_stats.cell(row=r_carr, column=5, value=coloc_en_car)
        c_pct = ws_stats.cell(row=r_carr, column=6, value=porcentaje)
        
        for c_obj in [c_nom, c_cnt, c_pct]:
            c_obj.fill = fill_zebra_even if r_carr % 2 == 0 else fill_zebra_odd
            c_obj.border = border_data
            c_obj.font = font_data
            
        c_cnt.alignment = align_center
        c_pct.alignment = align_center
        r_carr += 1

    ws_stats.column_dimensions["A"].width = 46
    ws_stats.column_dimensions["B"].width = 16
    ws_stats.column_dimensions["C"].width = 6
    ws_stats.column_dimensions["D"].width = 48
    ws_stats.column_dimensions["E"].width = 14
    ws_stats.column_dimensions["F"].width = 14

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
