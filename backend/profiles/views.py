import os
from rest_framework import viewsets
from rest_framework.permissions import IsAuthenticated
from django.http import FileResponse
from .models import Egresado, Empresa, PadronEgresado
from .serializers import (
    EgresadoSerializer, EmpresaSerializer, EmpresaStatusSerializer,
    PadronEgresadoSerializer
)
from .permissions import IsOwnerOrReadOnly
from rest_framework.response import Response
from rest_framework import status
from django.utils import timezone
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError, NotFound
import openpyxl

# Create your views here.

class EgresadoViewSet(viewsets.ModelViewSet):
    serializer_class = EgresadoSerializer
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]

    parser_classes = [MultiPartParser, FormParser, JSONParser]
    
    filterset_fields = ['carrera', 'es_verificado_padron', 'colocado', 'nivel_estudios']
    
    # Podemos buscar por matrícula, curp, o cruzar a la tabla User para buscar por nombre
    search_fields = [
        'matricula', 
        'curp', 
        'user__nombres', 
        'user__apellido_paterno', 
        'user__apellido_materno'
    ]
    
    ordering_fields = ['id', 'matricula']
    ordering = ['-id']

    def perform_create(self, serializer):
        user = self.request.user
        if user.rol not in ['egresado', 'soporte_ti', 'admin_uth'] and not user.is_superuser:
            raise PermissionDenied("Tu cuenta no tiene rol de Egresado para crear este perfil.")
        if hasattr(user, 'egresado'):
            raise ValidationError("Ya tienes un perfil de egresado asociado a tu cuenta.")
        
        # Validación de dos factores contra el Padrón Institucional:
        # Si la matrícula existe en el padrón, se valida automáticamente.
        # Si no existe, se permite el registro pero queda en espera de validación manual (es_verificado_padron=False).
        matricula = serializer.validated_data.get('matricula', '').strip()
        existe_en_padron = PadronEgresado.objects.filter(matricula__iexact=matricula).exists()
        
        serializer.save(user=user, es_verificado_padron=existe_en_padron)

    def get_queryset(self):
        return Egresado.objects.filter(user__deactivated_at__isnull=True)

    @action(detail=True, methods=['get'])
    def cv(self, request, pk=None):
        """Endpoint protegido para que el Frontend de Angular visualice o descargue el CV."""
        egresado = self.get_object()
        if not egresado.cv:
            raise NotFound("El egresado no cuenta con currículum vitae registrado.")
        
        user = request.user
        es_duenio = hasattr(user, 'egresado') and user.egresado == egresado
        es_admin = user.rol in ['admin_uth', 'soporte_ti'] or user.is_superuser
        # Una empresa solo puede ver el CV si el egresado tiene una postulación enviada a su empresa
        es_empresa_autorizada = (
            hasattr(user, 'empresa') and 
            egresado.postulaciones.filter(
                vacante__empresa=user.empresa,
                estado__in=['enviada_empresa', 'Aceptada', 'Rechazada']
            ).exists()
        )

        if not (es_duenio or es_admin or es_empresa_autorizada):
            raise PermissionDenied("No tienes autorización para consultar el currículum de este egresado.")

        try:
            archivo = egresado.cv.open('rb')
            response = FileResponse(archivo, content_type='application/pdf')
            nombre_archivo = os.path.basename(egresado.cv.name)
            response['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
            return response
        except FileNotFoundError:
            raise NotFound("El archivo físico del currículum no fue encontrado en el servidor.")

    @action(detail=True, methods=['get'])
    def documentos(self, request, pk=None):
        """Endpoint protegido para visualizar documentos anexos del egresado."""
        egresado = self.get_object()
        if not egresado.documentos:
            raise NotFound("El egresado no cuenta con documentos adicionales cargados.")

        user = request.user
        es_duenio = hasattr(user, 'egresado') and user.egresado == egresado
        es_admin = user.rol in ['admin_uth', 'soporte_ti'] or user.is_superuser

        if not (es_duenio or es_admin):
            raise PermissionDenied("Solo el propietario o administradores de la UTH pueden consultar los documentos de identidad.")

        try:
            archivo = egresado.documentos.open('rb')
            response = FileResponse(archivo, content_type='application/pdf')
            nombre_archivo = os.path.basename(egresado.documentos.name)
            response['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
            return response
        except FileNotFoundError:
            raise NotFound("El archivo físico de documentos no fue encontrado en el servidor.")

    @action(detail=True, methods=['patch'], url_path='verificar-padron')
    def verificar_padron(self, request, pk=None):
        """Permite al Administrador UTH validar o revocar manualmente la verificación en padrón."""
        if request.user.rol not in ['admin_uth', 'soporte_ti'] and not request.user.is_superuser:
            raise PermissionDenied("Solo administradores de la UTH pueden modificar la verificación de padrón.")

        egresado = self.get_object()
        nuevo_valor = request.data.get('es_verificado_padron', True)
        egresado.es_verificado_padron = bool(nuevo_valor)
        egresado.save()
        return Response(EgresadoSerializer(egresado).data, status=status.HTTP_200_OK)

    def destroy(self, request, *args, **kwargs):
        #SOFT DELETE USANDO DEACTIVATED_AT DE USER
        instance = self.get_object()

        # lógica de borrado suave
        instance.user.deactivated_at = timezone.now()
        instance.user.save()

        return Response(status=status.HTTP_204_NO_CONTENT)

class EmpresaViewSet(viewsets.ModelViewSet):
    serializer_class = EmpresaSerializer
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]

    parser_classes = [MultiPartParser, FormParser, JSONParser]
    
    filterset_fields = ['giro', 'sector', 'status']
    
    search_fields = [
        'nombre', 
        'actividad_de_la_empresa', 
        'correo_contacto'
    ]
    
    # Permitimos ordenar por status para que los admins vean rápido las "pendientes"
    ordering_fields = ['nombre', 'status', 'id']
    ordering = ['-id']

    def perform_create(self, serializer):
        if self.request.user.rol not in ['empresa', 'soporte_ti', 'admin_uth'] and not self.request.user.is_superuser:
            raise PermissionDenied("Tu cuenta no tiene rol de Empresa para crear este perfil.")
        if hasattr(self.request.user, 'empresa'):
            raise ValidationError("Ya tienes un perfil de empresa asociado a tu cuenta.")
        # Asignar el usuario autenticado al crear una empresa
        serializer.save(user=self.request.user)

    def get_queryset(self):
        return Empresa.objects.filter(user__deactivated_at__isnull=True)
    
    def destroy(self, request, *args, **kwargs):
        #SOFT DELETE USANDO DEACTIVATED_AT DE USER
        instance = self.get_object()

        # lógica de borrado suave
        instance.user.deactivated_at = timezone.now()
        instance.user.save()

        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(
        detail=True, 
        methods=['patch'], 
        serializer_class=EmpresaStatusSerializer
    )
    def cambiar_status(self, request, pk=None):
        #RBAC
        if request.user.rol not in ['admin_uth', 'soporte_ti'] and not request.user.is_superuser:
            raise PermissionDenied("No tienes permisos para aprobar o rechazar empresas.")

        # get empresa específica
        empresa = self.get_object()

        # inyectar los datos del request en el serializer
        serializer = self.get_serializer(empresa, data=request.data, partial=True)
        
        # validar y guardar los cambios
        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data, status=status.HTTP_200_OK)


class PadronEgresadoViewSet(viewsets.ModelViewSet):
    """
    Gestión del Padrón Oficial de Egresados UTH.
    Acceso restringido a Administradores UTH y Soporte TI.
    Permite la importación masiva desde archivos Excel (.xlsx) de 78 columnas.
    """
    queryset = PadronEgresado.objects.all()
    serializer_class = PadronEgresadoSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    
    filterset_fields = ['carrera', 'nivel', 'periodo', 'anio_egreso', 'estatus_titulacion']
    search_fields = ['matricula', 'nombre', 'curp', 'correo_personal', 'correo_escolares']
    ordering_fields = ['matricula', 'nombre', 'anio_egreso', 'fecha_importacion']
    ordering = ['-fecha_importacion', 'matricula']

    def check_admin_permission(self):
        user = self.request.user
        if user.rol not in ['admin_uth', 'soporte_ti'] and not user.is_superuser:
            raise PermissionDenied("Solo el personal administrativo de la UTH tiene acceso al Padrón Oficial.")

    def get_queryset(self):
        self.check_admin_permission()
        return self.queryset.all()

    def perform_create(self, serializer):
        self.check_admin_permission()
        serializer.save()

    def perform_update(self, serializer):
        self.check_admin_permission()
        serializer.save()

    def perform_destroy(self, instance):
        self.check_admin_permission()
        instance.delete()

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser])
    def importar(self, request):
        """
        Carga masiva de egresados desde archivo Excel (.xlsx).
        Lee las columnas del formato de Servicios Escolares / Encuesta de Seguimiento (78 columnas)
        y ejecuta un Upsert por matrícula.
        """
        self.check_admin_permission()

        archivo = request.FILES.get('archivo') or request.FILES.get('file')
        if not archivo:
            raise ValidationError({"archivo": "Debe proporcionar un archivo Excel (.xlsx)."})

        if not archivo.name.endswith(('.xlsx', '.xlsm')):
            raise ValidationError({"archivo": "Formato no válido. Debe ser un archivo .xlsx."})

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            sheet = wb['Hoja3'] if 'Hoja3' in wb.sheetnames else wb.active
        except Exception as e:
            raise ValidationError({"archivo": f"Error al abrir el archivo Excel: {str(e)}"})

        # Mapear encabezados de la fila 1
        headers = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val:
                headers[str(val).strip().lower()] = col_idx

        col_matricula = headers.get('matricula')
        if not col_matricula:
            raise ValidationError({"archivo": "El archivo no contiene la columna obligatoria 'Matricula'."})

        col_nombre = headers.get('nombre')
        col_carrera = headers.get('carrera')
        col_periodo = headers.get('periodo')
        col_anio = headers.get('año egreso') or headers.get('anio egreso')
        col_estatus = headers.get('estatus')
        col_estatus_tsu = headers.get('estatus tsu')
        col_etnia = headers.get('etnia indígena') or headers.get('etnia indigena')
        col_discapacidad = headers.get('discapacidad')
        col_genero = headers.get('género') or headers.get('genero')
        col_nivel = headers.get('nivel')
        col_tel_escolares = headers.get('tel escolares')
        col_correo_escolares = headers.get('correo escolares')
        col_domicilio = headers.get('domicilio')
        col_estado = headers.get('estado domicilio')
        col_municipio = headers.get('municipio')
        col_curp = headers.get('curp')
        col_fecha_nac = headers.get('fecha nacimiento')
        col_trabaja = headers.get('¿actualmente trabaja?') or headers.get('actualmente trabaja')
        col_correo_personal = headers.get('correo electrónico') or headers.get('correo electronico')
        col_tel_movil = headers.get('teléfono móvil') or headers.get('telefono movil')

        def get_val(row, col_idx):
            if not col_idx:
                return None
            v = sheet.cell(row=row, column=col_idx).value
            return str(v).strip() if v is not None else None

        creados = 0
        actualizados = 0
        omitidos = 0
        matriculas_procesadas = []

        for row_idx in range(2, sheet.max_row + 1):
            mat_raw = get_val(row_idx, col_matricula)
            if not mat_raw:
                omitidos += 1
                continue

            mat = mat_raw.split('.')[0] if '.' in mat_raw else mat_raw
            if not mat:
                omitidos += 1
                continue

            datos = {
                'nombre': get_val(row_idx, col_nombre) or "Sin Nombre",
                'carrera': get_val(row_idx, col_carrera) or "Sin Carrera",
                'periodo': get_val(row_idx, col_periodo),
                'anio_egreso': get_val(row_idx, col_anio),
                'estatus_titulacion': get_val(row_idx, col_estatus),
                'estatus_tsu': get_val(row_idx, col_estatus_tsu),
                'etnia_indigena': get_val(row_idx, col_etnia),
                'discapacidad': get_val(row_idx, col_discapacidad),
                'genero': get_val(row_idx, col_genero),
                'nivel': get_val(row_idx, col_nivel),
                'tel_escolares': get_val(row_idx, col_tel_escolares),
                'correo_escolares': get_val(row_idx, col_correo_escolares),
                'domicilio': get_val(row_idx, col_domicilio),
                'estado_domicilio': get_val(row_idx, col_estado),
                'municipio': get_val(row_idx, col_municipio),
                'curp': get_val(row_idx, col_curp),
                'fecha_nacimiento': get_val(row_idx, col_fecha_nac),
                'trabaja_actualmente': get_val(row_idx, col_trabaja),
                'correo_personal': get_val(row_idx, col_correo_personal),
                'telefono_movil': get_val(row_idx, col_tel_movil),
            }

            obj, created = PadronEgresado.objects.update_or_create(
                matricula=mat,
                defaults=datos
            )
            if created:
                creados += 1
            else:
                actualizados += 1
            matriculas_procesadas.append(mat)

        # Sincronizar automáticamente con las cuentas de egresados existentes
        if matriculas_procesadas:
            Egresado.objects.filter(matricula__in=matriculas_procesadas).update(es_verificado_padron=True)

        return Response({
            "mensaje": "Importación del Padrón completada exitosamente.",
            "total_filas_procesadas": len(matriculas_procesadas),
            "creados": creados,
            "actualizados": actualizados,
            "omitidos": omitidos
        }, status=status.HTTP_200_OK)